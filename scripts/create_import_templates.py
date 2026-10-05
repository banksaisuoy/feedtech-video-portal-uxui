import os
import csv
import shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Target directories
target_dirs = [
    r"C:\Users\nattakorn.sai\Downloads",
    r"C:\Users\nattakorn.sai\Downloads\Feedtech_Template",
    os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "public", "templates"))
]

for d in target_dirs:
    os.makedirs(d, exist_ok=True)

# Styling configuration
HEADER_FILL = PatternFill(start_color="006C49", end_color="006C49", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
REF_HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
REF_HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Segoe UI", size=10)
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

THIN_SIDE = Side(border_style="thin", color="CBD5E1")
BORDER_BOX = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

def style_worksheet(ws, header_fill=HEADER_FILL, header_font=HEADER_FONT):
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_BOX
    ws.row_dimensions[1].height = 28

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 22
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = DATA_FONT
            cell.border = BORDER_BOX
            cell.alignment = Alignment(vertical="center")
            if is_even and cell.fill.fill_type is None:
                cell.fill = ZEBRA_FILL

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

# ==========================================
# 1. USER IMPORT TEMPLATE (Strictly Matching Add User Modal)
# ==========================================
# Web Form Fields:
# 1. User ID (modalEmpId)
# 2. Full Name * (modalName)
# 3. Corporate Email * (modalEmail)
# 4. Department * (modalDept)
# 5. Role * (modalRoleSelect: User / Admin)
# 6. Executive Board Member (modalUserIsExecBoard: Yes / No)
# Removed: permission_level, is_admin, status (not in the web form)

user_headers = [
    "user_id",
    "full_name",
    "email",
    "department",
    "role",
    "executive_board"
]

user_samples = [
    ["EMP-1001", "Dr. Somchai Prasert", "somchai.p@feedtech.corp", "Biotech", "User", "Yes"],
    ["EMP-1002", "Ananya Srisuk", "ananya.s@feedtech.corp", "Swine", "User", "No"],
    ["EMP-1003", "Kittisak Tech", "kittisak.t@feedtech.corp", "Biotech", "Admin", "Yes"],
    ["EMP-1004", "Maria Wong", "maria.w@feedtech.corp", "QC-Lab", "User", "No"],
    ["EMP-1005", "David Miller", "david.m@feedtech.corp", "Raw Material", "User", "No"],
    ["EMP-1006", "Nattaporn Chaiyarat", "nattaporn.c@feedtech.corp", "Poultry", "User", "No"],
    ["EMP-1007", "Chen Wei", "chen.w@feedtech.corp", "China", "User", "No"],
    ["EMP-1008", "Supaporn Boonmee", "supaporn.b@feedtech.corp", "Aquatic", "User", "No"]
]

wb_user = Workbook()
ws_user = wb_user.active
ws_user.title = "Users_Import"
ws_user.append(user_headers)
for row in user_samples:
    ws_user.append(row)
style_worksheet(ws_user)

ws_user_ref = wb_user.create_sheet(title="Reference_Guides")
ws_user_ref.append(["Column / Field (ช่องในเว็บ)", "Allowed Values (ค่าที่ระบุได้)", "Required (จำเป็นไหม)", "Description / หมายเหตุ"])
ws_user_ref.append(["user_id (User ID)", "EMP-1001 หรือ 0001 (ตัวเลข/รหัสพนักงาน)", "ไม่บังคับ", "ถ้าเว้นว่างไว้ ระบบจะสุ่มเลข 4 หลักให้อัตโนมัติ"])
ws_user_ref.append(["full_name (Full Name *)", "ชื่อ - นามสกุล เช่น Dr. Somchai Prasert", "บังคับ (*)", "ชื่อจริงและนามสกุลของพนักงาน"])
ws_user_ref.append(["email (Corporate Email *)", "อีเมลองค์กร เช่น somchai.p@feedtech.corp", "บังคับ (*)", "อีเมลสำหรับใช้งานระบบ (ห้ามซ้ำกับผู้อื่น)"])
ws_user_ref.append(["department (Department *)", "Biotech, Swine, Aquatic, Conference, Dairy, Dairy Process, Extension Research, Nutrition, Oversea, Premix, Poultry, Raw Material, Ruminant, Ruminant Pakthongchai, Supplier, QC-Lab, China", "บังคับ (*)", "ต้องสะกดภาษาอังกฤษให้ตรงกับ 1 ใน 17 แผนกของระบบ"])
ws_user_ref.append(["role (Role *)", "User หรือ Admin", "บังคับ (*)", "User = พนักงานทั่วไป, Admin = ผู้ดูแลระบบ"])
ws_user_ref.append(["executive_board (Executive Board)", "Yes หรือ No", "ไม่บังคับ", "Yes = สมาชิกบอร์ดบริหาร (เข้าถึงได้ทุกคลิปแบบ VIP), No = ทั่วไป"])
style_worksheet(ws_user_ref, header_fill=REF_HEADER_FILL, header_font=REF_HEADER_FONT)

# ==========================================
# 2. VIDEO IMPORT TEMPLATE (Strictly Matching Upload Video Form)
# ==========================================
# Web Form Fields:
# 1. Video Title * (uploadVideoTitle)
# 2. Category * (uploadVideoDept)
# 3. Video Source Link * (uploadVideoUrl)
# 4. Description (uploadVideoDesc)
# 5. Access Control Policy * (uploadAccessMode: Public / Include / Exclude)
# 6. Allowed Personnel (allowed_personnel) - used when Access Policy is Include
# 7. Excluded Personnel (excluded_personnel) - used when Access Policy is Exclude
# Removed: duration, content_type, thumbnail_url (not in the upload input form)

video_headers = [
    "video_title",
    "category",
    "video_url",
    "description",
    "access_policy",
    "allowed_personnel",
    "excluded_personnel"
]

video_samples = [
    [
        "Industrial Bioreactor Fermentation Protocol",
        "Biotech",
        "/sample.mp4",
        "Standard operating procedure for industrial fermentation and microbial strain cultivation",
        "Public",
        "",
        ""
    ],
    [
        "Swine Climate Control & Automated Telemetry SOP",
        "Swine",
        "/sample.mp4",
        "Optimizing climate-controlled barns and automated ventilation telemetry for swine nurseries",
        "Include",
        "EMP-1001, EMP-1002, EMP-1003",
        ""
    ],
    [
        "Shrimp Biofloc RAS Water Quality Telemetry",
        "Aquatic",
        "/sample.mp4",
        "Real-time optical dissolved oxygen sensors and biofloc management in indoor aquaculture",
        "Public",
        "",
        ""
    ],
    [
        "Feed Mill Extruder Pellet Durability Index Compliance",
        "Raw Material",
        "/sample.mp4",
        "High-pressure steam extrusion maintenance and pellet durability testing index",
        "Exclude",
        "",
        "EMP-1005"
    ],
    [
        "Annual Animal Nutrition & Feed Formulation Townhall 2026",
        "Nutrition",
        "/sample.mp4",
        "Quarterly feed formulation trends, amino acid profiling, and sustainability milestones",
        "Public",
        "",
        ""
    ],
    [
        "Dairy Cold Chain & Milk Tank Microbial Safety SOP",
        "Dairy",
        "/sample.mp4",
        "Automated Clean-in-Place (CIP) sanitation protocols for bulk milk tanks and chiller lines",
        "Include",
        "EMP-1003, EMP-1004",
        ""
    ]
]

wb_video = Workbook()
ws_video = wb_video.active
ws_video.title = "Videos_Import"
ws_video.append(video_headers)
for row in video_samples:
    ws_video.append(row)
style_worksheet(ws_video)

ws_video_ref = wb_video.create_sheet(title="Reference_Guides")
ws_video_ref.append(["Column / Field (ช่องในเว็บ)", "Allowed Values (ค่าที่ระบุได้)", "Required (จำเป็นไหม)", "Description / หมายเหตุ"])
ws_video_ref.append(["video_title (Video Title *)", "ชื่อคลิปวิดีโอ เช่น Swine Immunity Assay Q4 Review", "บังคับ (*)", "หัวข้อหลักของวิดีโอ"])
ws_video_ref.append(["category (Category *)", "Biotech, Swine, Aquatic, Poultry, Dairy, Nutrition...", "บังคับ (*)", "หมวดหมู่แผนกเจ้าของเนื้อหา (ตรงกับ 1 ใน 17 แผนก)"])
ws_video_ref.append(["video_url (Video Source Link *)", "ลิงก์ OneDrive, SharePoint, หรือ /sample.mp4", "บังคับ (*)", "ที่อยู่ไฟล์วิดีโอสำหรับสตรีมมิ่ง"])
ws_video_ref.append(["description (Description)", "คำอธิบายเนื้อหาและบทสรุป", "ไม่บังคับ", "บทคัดย่อหรือหัวข้อย่อยของคลิป"])
ws_video_ref.append(["access_policy (Access Control Policy *)", "Public, Include, หรือ Exclude", "บังคับ (*)", "Public = ทุกคนดูได้, Include = ดูได้เฉพาะรายชื่อที่ระบุ, Exclude = ทุกคนดูได้ยกเว้นรายชื่อที่ระบุ"])
ws_video_ref.append(["allowed_personnel (Authorized Personnel)", "รหัสพนักงานคั่นด้วยลูกน้ำ เช่น EMP-1001, EMP-1002", "ใช้เมื่อเป็น Include", "รายชื่อพนักงานที่อนุญาตให้เปิดดูคลิปนี้ได้"])
ws_video_ref.append(["excluded_personnel (Excluded Personnel)", "รหัสพนักงานคั่นด้วยลูกน้ำ เช่น EMP-1005", "ใช้เมื่อเป็น Exclude", "รายชื่อพนักงานที่ห้ามเปิดดูคลิปนี้"])
style_worksheet(ws_video_ref, header_fill=REF_HEADER_FILL, header_font=REF_HEADER_FONT)

# Save files across all target directories
for d in target_dirs:
    # 1. User XLSX
    p_u_xlsx = os.path.join(d, "FeedTech_User_Import_Template.xlsx")
    wb_user.save(p_u_xlsx)

    # 2. User CSV (UTF-8 BOM)
    p_u_csv = os.path.join(d, "FeedTech_User_Import_Template.csv")
    with open(p_u_csv, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(user_headers)
        writer.writerows(user_samples)

    # 3. Video XLSX
    p_v_xlsx = os.path.join(d, "FeedTech_Video_Import_Template.xlsx")
    wb_video.save(p_v_xlsx)

    # 4. Video CSV (UTF-8 BOM)
    p_v_csv = os.path.join(d, "FeedTech_Video_Import_Template.csv")
    with open(p_v_csv, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(video_headers)
        writer.writerows(video_samples)

    print(f"Successfully updated templates in: {d}")
